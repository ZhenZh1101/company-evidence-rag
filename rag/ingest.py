"""Local, provenance-preserving readers for IR archives and ordinary documents."""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import subprocess
import tempfile
import warnings
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from .archive import normalize_record, read_manifest


@dataclass
class Source:
    key: str
    company: str
    title: str
    category: str
    publication_date: str | None
    publication_period: str | None
    date_basis: str
    scope: str
    source_url: str
    path: Path
    aliases: list[dict] = field(default_factory=list)


@dataclass
class Segment:
    text: str
    locator: str


SUPPORTED = {".html", ".htm", ".txt", ".md", ".pdf", ".docx", ".pptx", ".xlsx", ".xls", ".csv"}
NOTICE_FILES = {"streaming_links.txt", "online_viewers.txt", "online_links.txt", "unavailable.txt", "source_url.txt", "webcast.txt"}


def _safe_path(root: Path, relative: str) -> Path:
    candidate = (root / relative).resolve()
    if not candidate.is_relative_to(root.resolve()):
        raise ValueError(f"Path escapes source directory: {relative}")
    return candidate


def _source(path: Path, company: str, item: dict, meta: dict, url: str) -> Source:
    publication_date = item.get("publication_date", meta.get("publication_date"))
    if publication_date is not None:
        if not isinstance(publication_date, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", publication_date):
            raise ValueError(f"Publication date must be an ISO day or null: {publication_date!r}")
        date.fromisoformat(publication_date)
    return Source(
        key=hashlib.sha256(f"{company}\0{path}".encode()).hexdigest(), company=company,
        title=item.get("title") or meta.get("title") or path.stem,
        category=item.get("category") or meta.get("category") or "document",
        publication_date=publication_date,
        publication_period=item.get("publication_period", meta.get("publication_period")),
        date_basis=item.get("date_basis") or meta.get("date_basis") or "Unknown publication date",
        scope=item.get("scope") or meta.get("scope") or "local", source_url=url, path=path,
    )


def _selection_reason(path: str, candidates: set[str], category: str, meta: dict, file_record: dict | None = None) -> str | None:
    """Suppress only known archive wrappers and alternative representations."""
    p = Path(path)
    if p.suffix.lower() not in SUPPORTED:
        return "Unsupported format or non-document asset"
    if p.name in NOTICE_FILES or p.name.endswith(".excerpt.txt") or p.name.startswith("unavailable_") or "streaming" in p.name or "online_viewer" in p.name:
        return "Media link/availability metadata; no transcript content"
    if category.startswith("sec_"):
        primary = Path(meta.get("sec_submission", {}).get("primaryDocument", "")).name
        primary_paths = {primary, f"sec_documents/{primary}"} & candidates
        primary_html = bool(primary_paths) and Path(primary).suffix.lower() in {".html", ".htm"}
        rendered_primary = any(f"primary_document{ext}" in candidates for ext in (".html", ".pdf", ".txt"))
        accession = meta.get("accession", "")
        if p.name.endswith(("-index.html", "-index-headers.html")) or p.name in {
            "page.html", "page.txt", "article.html", "ir_filing.html", "sec_filing_index.html",
            "filing_detail.html", "filing_detail.txt", "source_listing_row.html",
        }:
            return "SEC filing directory/detail wrapper"
        if p.name == "complete_submission.txt" or (accession and p.name == f"{accession}.txt"):
            return "SEC submission bundle duplicates filings and can contain encoded binary assets"
        if re.fullmatch(r"R\d+\.htm", p.name) and primary_html:
            return "SEC generated financial-table view already contained in primary filing"
        if p.name in {"filing.html", "html.html", "primary_document.html"} and primary_html:
            return "IR HTML mirror of SEC primary filing"
        if p.name in {"ir_filing.pdf", "ir_filing.docx", "ir_filing.xlsx"} and (primary_paths or rendered_primary):
            return "Alternative format of SEC primary filing"
        # The IR accession-named formats transform the same filing. Keep exhibits.
        versions = [x for x in candidates if accession and Path(x).name in {
            f"{accession}.pdf", f"{accession}.docx", f"{accession}.rtf.docx", f"{accession}.xls", f"{accession}.xlsx"}]
        transformed = path in versions
        if transformed:
            if primary_html:
                return "Alternative format of SEC primary HTML filing"
            versions.sort(key=lambda x: ([".pdf", ".docx", ".xls", ".xlsx"].index(Path(x).suffix.lower()), x))
            if versions and path != versions[0]:
                return "Alternative format of selected SEC filing"
        if p.name in {"filing.html", "html.html"} and not primary_html and any(Path(x).suffix.lower() == ".pdf" for x in versions):
            return "IR filing mirror; selected rendered PDF retains ownership-form labels"
    else:
        record = file_record or {}
        material_url = meta.get("url")
        if ((category in {"financial_results", "annual_reports"} and p.name in {"page.html", "page.txt", "article.html"})
                or (category in {"quarterly_results", "annual_report", "investor_presentation", "current_governance_document"}
                    and p.name in {"source.html", "source.txt"})
                or re.fullmatch(r"source_page_\d+\.html", p.name)):
            return "Parent collection/homepage capture; snapshot is not dated document evidence"
        attachment_record = any(isinstance(f, dict) and f.get("source_url") == material_url
                                and f.get("kind") in {"attachment", "document_attachment", "pdf_attachment", "alternative_attachment"}
                                for f in meta.get("files", []))
        if (p.name in {"page.html", "page.txt", "article.html"}
                and record.get("kind") in {"source_html", "rendered_html", "readable_text", "original_http_html",
                                           "readable_offline_html", "page_text", "original_html", "article_text", "readable_article_html"}
                and record.get("source_url") and material_url and record["source_url"] != material_url
                and (attachment_record or category == "quarterly_results")):
            return "Parent collection/homepage capture; snapshot is not dated document evidence"
        preference = ("article.html", "page.html", "page.txt") if category in {"ir_news", "press_release"} else ("article.html", "page.txt", "page.html")
        readable = next((x for x in preference if x in candidates), None)
        if path in {"article.html", "page.txt", "page.html"} and path != readable:
            return "Alternative webpage representation"
        if p.name in {"content.txt", "source.txt"} and "source.html" in candidates:
            return "Alternative webpage representation"
        localized_text = re.fullmatch(r"content_([\w-]+)\.txt", p.name)
        if localized_text and f"page_{localized_text[1]}.html" in candidates:
            return "Alternative webpage representation"
        if readable and p.name in {"source_listing.html", "source_fragment.html", "listing_entry.html", "listing.txt", "ir_page.html", "ir_page.txt", "ir_article.html"}:
            return "Archive listing/capture/mirror already represented by canonical article"
        if readable and (p.name.startswith("main_site_") or p.name.startswith("ir_mirror_")):
            return "Merged release mirror"
        if p.name == "release.pdf" and category in {"ir_news", "press_releases", "news"} and readable:
            return "PDF version of the same complete press release"
    return None


def discover(root: Path, company: str | None = None) -> tuple[list[Source], list[dict]]:
    """Read final archive index + registered files; never crawl audit/ as evidence."""
    root = root.expanduser().resolve()
    if not root.exists():
        raise FileNotFoundError(root)
    report: list[dict] = []
    sources: list[Source] = []
    records = read_manifest(root) if root.is_dir() else None
    if records is not None:
        for item in records:
            try:
                if not isinstance(item, dict) or not item.get("folder"):
                    raise ValueError("Archive index entry has no folder")
                folder = _safe_path(root, item["folder"])
                meta_path = _safe_path(root, item["meta_path"]) if item.get("meta_path") else _safe_path(folder, "meta.json")
                meta = json.loads(meta_path.read_text())
                if not isinstance(meta, dict):
                    raise ValueError("Archive metadata must be an object")
                item, meta = normalize_record(item, meta)
                if item.get("status") == "unavailable" or meta.get("status") == "unavailable":
                    report.append({"path": str(folder), "reason": "Unavailable source; metadata only"})
                    continue
                files = meta.get("files")
                if not isinstance(files, list):
                    raise ValueError("Archive metadata files must be a list")
                candidates = {f.get("path") or f.get("filename") for f in files if isinstance(f, dict)} - {None}
                for f in files:
                    if not isinstance(f, dict):
                        report.append({"path": str(folder), "reason": "Invalid file record", "kind": "error"})
                        continue
                    relative = f.get("path") or f.get("filename")
                    if not relative:
                        report.append({"path": str(folder), "reason": "File record has no path", "kind": "error"})
                        continue
                    try:
                        path = _safe_path(folder, relative)
                        if f.get("status") in {"failed", "unavailable"} or not path.is_file():
                            raise ValueError("File unavailable or download failed")
                        reason = _selection_reason(relative, candidates, item.get("category", ""), meta, f)
                        if reason:
                            report.append({"path": str(path), "reason": reason})
                            continue
                        issuer = company or meta.get("ticker") or root.name.split("_")[0]
                        source_item = item
                        url = f.get("source_url") or item.get("url") or meta.get("url", "")
                        if "publication_date" in f or "publication_period" in f:
                            source_item = {**item, "publication_date": f.get("publication_date"),
                                           "publication_period": f.get("publication_period"),
                                           "date_basis": f.get("date_basis") or "Attachment date supplied by archive"}
                        elif (issuer == "NOK" and f.get("kind") in {"attachment", "document_attachment", "pdf_attachment", "alternative_attachment"}
                              and item.get("category", "").startswith(("blog", "editorial", "corporate_site", "corporate_event", "ir_event", "technology"))
                              and url != (item.get("url") or meta.get("url"))):
                            source_item = {**item, "publication_date": None, "publication_period": None,
                                           "date_basis": "Attachment publication date unverified; parent page date does not date linked documents"}
                        source = _source(path, issuer, source_item, meta, url)
                        # Preserve release URL separately when an attachment has its own URL.
                        if item.get("url") and item["url"] != source.source_url:
                            source.aliases.append({"source_url": item["url"], "title": source.title,
                                                   "publication_date": item.get("publication_date", meta.get("publication_date")),
                                                   "publication_period": item.get("publication_period", meta.get("publication_period")),
                                                   "date_basis": item.get("date_basis") or meta.get("date_basis") or source.date_basis, "scope": source.scope,
                                                   "role": "release_page"})
                        sources.append(source)
                    except (ValueError, OSError, TypeError) as exc:
                        report.append({"path": str(folder / str(relative)), "reason": str(exc), "kind": "error"})
            except (ValueError, OSError, KeyError, TypeError) as exc:
                report.append({"path": str(item.get("folder", root) if isinstance(item, dict) else root), "reason": str(exc), "kind": "error"})
        return sources, report
    base = root if root.is_dir() else root.parent
    paths = sorted(root.rglob("*")) if root.is_dir() else [root]
    for path in paths:
        if not path.is_file() or any(part.startswith(".") or part in {"audit", "__pycache__", "node_modules"} for part in path.relative_to(base).parts):
            continue
        if not path.resolve().is_relative_to(base):
            report.append({"path": str(path), "reason": "Symlink escapes source directory", "kind": "error"})
        elif path.suffix.lower() in SUPPORTED:
            sources.append(_source(path.resolve(), company or base.name, {}, {}, ""))
        else:
            report.append({"path": str(path), "reason": "Unsupported format"})
    return sources, report


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\xa0", " ").replace("\r", "")
    return re.sub(r"\n[ \t]*\n(?:[ \t]*\n)+", "\n\n", text).strip()


def _table_text(rows: list[list[str]], context: str = "") -> str:
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        return ""
    # Keep the first heading rows with each later table chunk; years are headers.
    heading_count = 1
    for row in rows[1:5]:
        values = [v.strip().replace(",", "").replace("$", "").replace("%", "") for v in row]
        numbers = [v for v in values if re.fullmatch(r"\(?-?\d+(?:\.\d+)?\)?", v)]
        if numbers and any(not re.fullmatch(r"20\d{2}", v) for v in numbers):
            break
        heading_count += 1
    header = "\n".join(" | ".join(row) for row in rows[:heading_count])
    body = "\n".join(" | ".join(row) for row in rows[heading_count:])
    return _clean((f"Table context: {context}\n" if context else "") + header + "\n[Rows]\n" + body)


def _html_table(table) -> list[list[str]]:
    rows, pending = [], {}
    for tr in table.find_all("tr"):
        if tr.find_parent("table") is not table:
            continue
        cells, column = {}, 0
        for col, (value, remaining) in list(pending.items()):
            cells[col] = value
            if remaining <= 1:
                del pending[col]
            else:
                pending[col] = (value, remaining - 1)
        for cell in tr.find_all(["td", "th"], recursive=False):
            while column in cells:
                column += 1
            value = " ".join(cell.stripped_strings)
            try:
                colspan, rowspan = min(64, max(1, int(cell.get("colspan", 1)))), min(1024, max(1, int(cell.get("rowspan", 1))))
            except (TypeError, ValueError):
                colspan = rowspan = 1
            for col in range(column, column + colspan):
                cells[col] = value if col == column else ""
                if rowspan > 1:
                    pending[col] = (value if col == column else "", rowspan - 1)
            column += colspan
        if cells:
            rows.append([cells.get(col, "") for col in range(max(cells) + 1)])
    # SEC HTML uses many spacer columns. Remove only columns empty in every row.
    width = max((len(row) for row in rows), default=0)
    keep = [col for col in range(width) if any(col < len(row) and row[col].strip() for row in rows)]
    return [[row[col] if col < len(row) else "" for col in keep] for row in rows]


def _table_context(table) -> str:
    previous = table.find_previous(["h1", "h2", "h3", "p"])
    semantic = " ".join(previous.stripped_strings)[:500] if previous else ""
    captions, remaining, inspected, boundary = [], 600, 0, False
    node = table
    # Caption DIVs often precede the table's wrapper, not the table itself.
    for _ in range(4):  # Table plus at most three parent wrappers.
        for sibling in node.previous_siblings:
            inspected += 1
            if inspected > 40:
                break
            if (getattr(sibling, "name", None) in {"table", "hr"}
                    or (hasattr(sibling, "find_all") and sibling.find(["table", "hr"]))):
                boundary = True
                break
            text = " ".join(sibling.stripped_strings) if hasattr(sibling, "stripped_strings") else str(sibling).strip()
            if text:
                # Walking backwards: retain the text closest to this table.
                piece = text[-remaining:]
                captions.append(piece)
                remaining -= len(piece) + 1
                if remaining <= 0:
                    break
        if boundary or remaining <= 0 or inspected > 40:
            break
        node = node.parent
        if node is None or node.name in {"body", "html", "[document]"}:
            break
    nearby = "\n".join(reversed(captions))
    unit_pattern = re.compile(r"\b(?:in\s+(?:thousands|millions|billions)|per[ -]+share(?:\s+amounts)?|(?:U\.?S\.?\s*)?dollars)\b", re.I)
    units = lambda text: {re.sub(r"\s+", " ", match.group().lower()) for match in unit_pattern.finditer(text)}
    if units(nearby) and (units(nearby) - units(semantic) or (boundary and semantic not in nearby)):
        return nearby
    if units(semantic) and boundary and semantic not in nearby:
        return ""  # Do not borrow a unit caption from a preceding table.
    return semantic


def _html(path: Path) -> list[Segment]:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(path.read_bytes(), "html.parser")
    hidden_style = re.compile(r"(?:^|;)\s*(?:display\s*:\s*none|visibility\s*:\s*(?:hidden|collapse))\s*(?:!important\s*)?(?:;|$)", re.I)
    noncontent_tags = {"script", "style", "nav", "header", "footer", "noscript", "svg", "form",
                       "ix:hidden", "ix:header", "ix:resources", "ix:references", "link:schemaref", "link:linkbaseref"}
    for tag in list(soup.find_all()):
        # A parent's removal also decomposes its descendants in this snapshot.
        if tag.name is None or tag.attrs is None:
            continue
        if (tag.name in noncontent_tags or tag.name.startswith(("xbrli:", "xbrldi:"))
                or tag.has_attr("hidden") or str(tag.get("aria-hidden", "")).lower() == "true"
                or hidden_style.search(tag.get("style", ""))):
            tag.decompose()
    main = soup.find("article") or soup.find("main") or soup.body or soup
    segments, section, paragraphs = [], "document", []

    def flush():
        if paragraphs:
            segments.append(Segment(_clean("\n\n".join(paragraphs)), section))
            paragraphs.clear()

    # SEC filings often use nested div/span instead of semantic paragraphs.
    table_segments = []
    # Extract children before removing layout parents, preserving actual nested tables.
    for table_number, table in reversed(list(enumerate(main.find_all("table"), 1))):
        text = _table_text(_html_table(table), _table_context(table))
        if text:
            table_segments.append(Segment(text, f"table {table_number}"))
        table.decompose()
    segments.extend(reversed(table_segments))
    for tag in main.find_all(["h1", "h2", "h3", "h4", "p", "li"]):
        if tag.find_parent(["p", "li"]):
            continue
        text = " ".join(tag.stripped_strings)
        if not text:
            continue
        if tag.name.startswith("h"):
            flush()
            section = text[:180]
        paragraphs.append(text)
    flush()
    # Use all remaining visible text for div-heavy filings, which otherwise lose prose.
    if not segments or (len(" ".join(s.text for s in segments if not s.locator.startswith("table "))) < len(main.get_text(" ", strip=True)) * .65):
        prose = _clean(main.get_text("\n", strip=True))
        segments = [s for s in segments if s.locator.startswith("table ")]
        if prose:
            segments.append(Segment(prose, "document"))
    return segments


def _ocr_pdf_page(path: Path, number: int) -> str:
    with tempfile.TemporaryDirectory(prefix="rag-ocr-") as directory:
        image = Path(directory) / "page"
        commands = [
            ["pdftoppm", "-f", str(number), "-l", str(number), "-singlefile", "-r", "180", "-png", str(path), str(image)],
            ["tesseract", str(image.with_suffix(".png")), "stdout", "-l", "eng"],
        ]
        for command in commands:
            result = subprocess.run(command, capture_output=True, text=True, timeout=120)
            if result.returncode:
                raise ValueError(f"{command[0]} failed: {result.stderr[:500]}")
        return result.stdout


def extract(source: Source, ocr: bool = False) -> list[Segment]:
    """Extract evidence, raising a visible error if the input has no readable text."""
    path, suffix = source.path, source.path.suffix.lower()
    segments: list[Segment] = []
    if suffix in {".html", ".htm"}:
        segments = _html(path)
    elif suffix in {".txt", ".md"}:
        segments = [Segment(path.read_text(encoding="utf-8-sig", errors="replace"), "text")]
    elif suffix == ".pdf":
        if shutil.which("pdftotext"):
            result = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", str(path), "-"],
                                    capture_output=True, text=True, timeout=120)
            if result.returncode:
                raise ValueError(f"PDF extraction failed for {path}: {result.stderr[:500]}")
            pages = result.stdout.split("\f")
            if pages and not pages[-1].strip():
                pages.pop()
        else:
            from pypdf import PdfReader
            pages = [page.extract_text(extraction_mode="layout") or "" for page in PdfReader(path).pages]
        if ocr and any(not page.strip() for page in pages):
            missing = [name for name in ("pdftoppm", "tesseract") if not shutil.which(name)]
            if missing:
                raise ValueError(f"Local PDF OCR requires {', '.join(missing)} on PATH; install Poppler/Tesseract with English language data and rerun with --ocr")
        for number, text in enumerate(pages, 1):
            if text.strip():
                segments.append(Segment(text, f"page {number}"))
            elif ocr:
                try:
                    recognized = _ocr_pdf_page(path, number)
                    if recognized.strip():
                        segments.append(Segment(recognized, f"page {number} (OCR; verify against original)"))
                        warnings.warn(f"{path}: page {number} used local English OCR; verify against original", stacklevel=2)
                    else:
                        warnings.warn(f"{path}: page {number} OCR returned no text; inspect the original image", stacklevel=2)
                except (ValueError, OSError, subprocess.TimeoutExpired) as exc:
                    warnings.warn(f"{path}: page {number} OCR failed: {exc}; inspect original scan and Tesseract language data", stacklevel=2)
            else:
                warnings.warn(f"{path}: page {number} has no extractable text; image/OCR review required", stacklevel=2)
    elif suffix == ".docx":
        from docx import Document
        document = Document(path)
        text = "\n\n".join(p.text for p in document.paragraphs if p.text.strip())
        if text:
            segments.append(Segment(text, "document paragraphs"))
        for number, table in enumerate(document.tables, 1):
            segments.append(Segment(_table_text([[c.text for c in row.cells] for row in table.rows]), f"table {number}"))
    elif suffix == ".pptx":
        from pptx import Presentation
        for number, slide in enumerate(Presentation(path).slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.append(_table_text([[c.text for c in row.cells] for row in shape.table.rows]))
            if texts:
                segments.append(Segment("\n\n".join(texts), f"slide {number}"))
            else:
                warnings.warn(f"{path}: slide {number} has no extractable text; image review required", stacklevel=2)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                rows = [["" if c is None else str(c) for c in row] for row in sheet.iter_rows(values_only=True)]
                segments.append(Segment(_table_text(rows), f"sheet {sheet.title}"))
        finally:
            workbook.close()
    elif suffix == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(path)
        for sheet in workbook.sheets():
            rows = [[str(c.value) for c in sheet.row(i)] for i in range(sheet.nrows)]
            segments.append(Segment(_table_text(rows), f"sheet {sheet.name}"))
    elif suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            segments = [Segment(_table_text(list(csv.reader(handle))), "table 1")]
    else:
        raise ValueError(f"Unsupported document format: {suffix}")
    result = [Segment(_clean(s.text), s.locator) for s in segments if _clean(s.text)]
    if not result:
        if suffix == ".pdf" and ocr:
            raise ValueError(f"No extractable text in {path}; OCR produced no usable text, inspect scan quality and Tesseract language data")
        raise ValueError(f"No extractable text in {path}; scanned/image-only documents require OCR")
    return result


def chunk_segments(segments: list[Segment], max_chars: int = 2600, overlap: int = 250) -> list[Segment]:
    """Keep source boundaries and repeat table headers when splitting long tables."""
    if max_chars < 200 or not 0 <= overlap < max_chars // 2:
        raise ValueError("max_chars must be >= 200 and overlap in [0, max_chars / 2)")
    chunks: list[Segment] = []
    for segment in segments:
        text = _clean(segment.text)
        if not text:
            continue
        prefix = ""
        if "\n[Rows]\n" in text:
            header, body = text.split("\n[Rows]\n", 1)
            prefix = header[:max_chars // 3] + "\n[Rows]\n"
            # Long headings must still be indexed in full, even if only their start repeats.
            text = body if len(header) <= max_chars // 3 else text
        elif segment.locator.startswith("page ") and len(text) > max_chars:
            # PDF layout extraction preserves rows; repeat page/column/units context.
            lines = text.splitlines()
            prefix = "\n".join(lines[:min(8, len(lines))])[:max_chars // 4] + "\n"
        capacity = max_chars - len(prefix)
        start, part = 0, 1
        while start < len(text):
            end = min(start + capacity, len(text))
            if end < len(text):
                split = text.rfind("\n", start + capacity // 2, end)
                if split < 0:
                    split = text.rfind(" ", start + capacity // 2, end)
                if split >= 0:
                    end = split
            body = text[start:end].strip()
            if body:
                chunks.append(Segment(prefix + body, f"{segment.locator}, part {part}"))
                part += 1
            if end == len(text):
                break
            # Table rows stay intact. For prose, overlap only at a word boundary.
            next_start = end if prefix else max(start + 1, end - overlap)
            if not prefix and next_start > 0 and not text[next_start - 1].isspace():
                boundary = text.find(" ", next_start, end)
                next_start = boundary + 1 if boundary >= 0 else end
            start = next_start
        if not text and prefix.strip():
            chunks.append(Segment(prefix, segment.locator))
    return chunks
